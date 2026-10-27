"""Create the editable Word working copy from the shared plan-data.json.

Run only when creating a new Word baseline. Preserve any team edits in an
existing working copy before regenerating it.
"""
from pathlib import Path
import json
from urllib.parse import urljoin
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'plan-data.json').read_text())
OUT = ROOT / ('Elite-Fit-Internal-Project-Plan-v' + DATA['metadata']['version'] + '.docx')
BASE = 'https://digitalmovementuk.github.io/elitefit-ops/project-management/internal-plan/'
INK, SAND, CORAL, PAPER, MUTED = '1D2123', 'DFDDD0', 'FF694B', 'F6F5F0', '53564F'
ROLES = {r['id']: r for r in DATA['roles']}
PHASES = {p['id']: p for p in DATA['phases']}
ACTS = {a['id']: (r, a) for r in DATA['schedule'] for a in r['activities']}

doc = Document()
doc.core_properties.title = DATA['metadata']['title']
doc.core_properties.subject = 'Editable internal project plan — version ' + DATA['metadata']['version']
doc.core_properties.author = 'Digital Movement UK'
doc.core_properties.keywords = 'EliteFit, implementation, editable project plan'
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title', 29), ('Heading 1', 23), ('Heading 2', 14), ('Heading 3', 11)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(INK)
    style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = True
doc.styles['Heading 2'].font.color.rgb = RGBColor.from_string('A5321D')
doc.styles['List Bullet'].paragraph_format.space_after = Pt(4)


def section_setup(s, landscape=False):
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
    s.top_margin, s.bottom_margin = Mm(17), Mm(17)
    s.left_margin, s.right_margin = Mm(18), Mm(18)
    s.header_distance, s.footer_distance = Mm(8), Mm(8)


section_setup(doc.sections[0])
header = doc.sections[0].header.paragraphs[0]
header.text = 'ELITEFIT   /   INTERNAL WORKING PLAN'
header.runs[0].font.size = Pt(8)
header.runs[0].font.bold = True
header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
footer = doc.sections[0].footer.paragraphs[0]
footer.text = 'v' + DATA['metadata']['version'] + ' · Pre-planning · Indicative months only     |     '
footer.runs[0].font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def shade(cell, color):
    pr = cell._tc.get_or_add_tcPr()
    for existing in pr.findall(qn('w:shd')):
        pr.remove(existing)
    node = OxmlElement('w:shd')
    node.set(qn('w:fill'), color)
    pr.append(node)


def no_split(row):
    node = OxmlElement('w:cantSplit')
    row._tr.get_or_add_trPr().append(node)


def table(headers, rows, widths, size=9.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    preferred = t._tbl.tblPr.find(qn('w:tblW'))
    preferred.set(qn('w:type'), 'dxa')
    preferred.set(qn('w:w'), str(round(sum(widths) / 25.4 * 1440)))
    for c, width in zip(t.columns, widths):
        c.width = Mm(width)
    for row_values in [headers] + list(rows):
        row = t.rows[0] if row_values is headers else t.add_row()
        no_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.width = Mm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                v = OxmlElement('w:' + side)
                v.set(qn('w:w'), '85')
                v.set(qn('w:type'), 'dxa')
                margins.append(v)
            pr.append(margins)
            shade(cell, INK if row_values is headers else (PAPER if len(t.rows) % 2 == 0 else 'FFFFFF'))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.07
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.bold = row_values is headers
                    if row_values is headers:
                        run.font.color.rgb = RGBColor.from_string('FFFFFF')
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def p(text='', bold=False, small=False):
    x = doc.add_paragraph()
    r = x.add_run(text)
    r.bold = bold
    if small:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(MUTED)
    return x


def bullet(text):
    doc.add_paragraph(text, 'List Bullet')


def link(label, url):
    x = doc.add_paragraph()
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), x.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), 'A5321D')
    rp.append(color)
    run.append(rp)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    h.append(run)
    x._p.append(h)


def heading(label, subtitle='', new=True):
    h = doc.add_paragraph(label, 'Heading 1')
    h.paragraph_format.page_break_before = new
    if subtitle:
        p(subtitle)


def notes(text):
    doc.add_paragraph('Working notes', 'Heading 3')
    p(text, small=True)


def active(row, phase):
    return [a for a in row['activities'] if phase in a['phase_ids']]


def row_for(task):
    return ACTS[task['activity_ids'][0]][0]


def task_phases(task):
    ids = {phase for aid in task['activity_ids'] for phase in ACTS[aid][1]['phase_ids']}
    return [PHASES[phase]['label'] for phase in PHASES if phase in ids]


# The seven month views, Gantt and task ownership use the same schedule objects.
p('VERSION ' + DATA['metadata']['version'] + '   ·   EDITABLE TEAM COPY', bold=True, small=True)
doc.add_paragraph('EliteFit\nInternal project plan', 'Title')
p('Website Months 1–4 · Full integration Months 5–7 · SEO Months 5–12', bold=True)
p('Project start and exact dates are still to be agreed. The technical integration lead still needs to be named.')
p('How to edit: change the text and tables directly. Use Comments or Review → Track Changes for proposed changes. Betzy records agreed changes in the shared plan so this copy, the online timeline and the task list stay aligned.', small=True)
link('Open the online plan', BASE)
first = DATA['slides'][0]
doc.add_paragraph('1. Roles and deliverables', 'Heading 1')
table(['Person / role', 'Responsibility', 'Deliverable'], [
    [r['name'] + '\n' + r['role'], r['responsibility'], r['deliverable']] for r in DATA['roles']
], [44, 68, 62], size=8.7)
notes(first['note'])

s = doc.add_section(WD_SECTION_START.NEW_PAGE)
section_setup(s)
g = DATA['slides'][1]
heading('2. One shared Gantt chart', g['subtitle'], new=False)
p('Each coloured cell is planned work. The words inside the cells can be edited.', small=True)
gr = []
for row in DATA['schedule']:
    name = ROLES[row['owner_id']]['name']
    if row['owner_id'] == 'technical':
        name = 'Technical lead · to confirm'
    gr.append([row['title'] + '\n' + name] + [' / '.join(a['short_label'] for a in active(row, phase)) or '—' for phase in PHASES])
t = table(['Workstream / responsible'] + [v['short_label'] for v in PHASES.values()], gr, [46] + [16] * 8, size=7.7)
for row in t.rows[1:]:
    for cell in row.cells[1:]:
        if cell.text != '—':
            shade(cell, 'FBC6BB')
p('First 20 SEO pages across Months 5–7 after website approval. SEO then continues through Month 12 at the agreed volume.', bold=True)
notes(g['note'])

s = doc.add_section(WD_SECTION_START.NEW_PAGE)
section_setup(s)
for i, slide in enumerate([s for s in DATA['slides'] if s['kind'] == 'month']):
    phase = slide['phase_id']
    heading(PHASES[phase]['label'] + ' — ' + slide['title'], slide['subtitle'], new=i > 0)
    rows = []
    for row in DATA['schedule']:
        acts = active(row, phase)
        if acts:
            rows.append([row['title'] + '\n' + ROLES[row['owner_id']]['name'], '\n'.join(a['label'] for a in acts), '\n'.join(a['deliverable'] for a in acts)])
    table(['Work / responsible', 'Key tasks', 'Deliverable'], rows, [47, 56, 71], size=9.4)
    notes(slide['note'])

slide = next(s for s in DATA['slides'] if s['kind'] == 'integration')
heading('3. Full technical integration', slide['subtitle'])
x = DATA['integration']
p(x['delivery'], bold=True)
doc.add_paragraph('Known now', 'Heading 2')
table(['Finding', 'Current understanding'], [[r['title'], r['text']] for r in x['known']], [51, 123])
doc.add_paragraph('Check before build', 'Heading 2')
for text in x['check']:
    bullet(text)
doc.add_paragraph('Journey to connect', 'Heading 2')
p(' → '.join(x['sequence']), bold=True)
notes(slide['note'])

heading('4. Backlog overview', '24 planned work packages, prepared for Jira. Completion has not been checked.')
table(['Workstream', 'Responsible', 'Packages'], [[r['title'], ROLES[r['owner_id']]['name'], ', '.join(t['id'] for t in DATA['tasks'] if row_for(t)['id'] == r['id'])] for r in DATA['schedule']], [60, 65, 49])
p('The full editable task list follows the communication guide. No Jira issues have been created.', small=True)
notes(next(s for s in DATA['slides'] if s['kind'] == 'backlog')['note'])

heading('5. Betzy’s communication guide', next(s for s in DATA['slides'] if s['kind'] == 'voice')['subtitle'])
for principle in DATA['voice']['principles']:
    doc.add_paragraph(principle['title'], 'Heading 2')
    p(principle['text'])
p(DATA['voice']['reliability'], bold=True)
for example in DATA['voice']['examples'][:4]:
    doc.add_paragraph(example['label'], 'Heading 3')
    p(example['text'])
heading('Client reply examples', 'Use a specific date, time and time zone for every promised update.')
for example in DATA['voice']['examples'][4:]:
    doc.add_paragraph(example['label'], 'Heading 2')
    p(example['text'])
notes(next(s for s in DATA['slides'] if s['kind'] == 'voice')['note'])

for i, task in enumerate(DATA['tasks']):
    if i % 2 == 0:
        heading('6. Full task backlog' if i == 0 else 'Task backlog — continued', 'Planned work · completion not checked')
    doc.add_paragraph(task['id'] + ' — ' + task['task'], 'Heading 2')
    r = row_for(task)
    table(['Responsible', 'Timing', 'Status'], [[ROLES[r['owner_id']]['name'], ', '.join(task_phases(task)), task['status']]], [55, 64, 55], size=9)
    p('Result / acceptance: ' + task['outcome'])
    p('Helpers / approval: ' + task['helpers'], small=True)
    p('Needs: ' + (', '.join(task['dependencies']) or 'None'), small=True)
    if task.get('dependency_note'):
        p(task['dependency_note'], small=True)
    for text in task.get('checks', []):
        bullet(text)
    p('Team notes / next step: ___________________________________', small=True)

heading('7. Open decisions and approvals', 'Record decisions here, then keep the shared plan up to date.')
table(['Decision', 'Owner', 'Status / note'], [[d['question'], ROLES[d['owner_id']]['name'], d['status'] + ' — ' + d['note']] for d in DATA['decisions']], [73, 28, 73], size=9)
doc.add_paragraph('Required approvals', 'Heading 2')
for rule in DATA['approval_rules']:
    bullet(rule)

heading('8. Sources and shared files', 'These links point to the existing project references. No new research is claimed in this Word conversion.')
for source in DATA['sources']:
    link(source['title'], urljoin(BASE, source['url']))
doc.add_paragraph('Additional integration evidence', 'Heading 2')
for source in DATA['integration'].get('sources', []):
    p(source['title'], bold=True)
    p(source['detail'], small=True)
    if source.get('path'):
        link('Open operational reference', urljoin(BASE, source['path']))
    elif source.get('availability'):
        p(source['availability'], small=True)
doc.add_paragraph('Keeping one agreed plan', 'Heading 2')
p('The Word file is an editable team copy of version ' + DATA['metadata']['version'] + '. Betzy collects team changes and updates plan-data.json for agreed changes. The online plan, Gantt, Markdown and task CSV are then regenerated from that shared source.')
link('Shared plan source', BASE + 'plan-data.json')
p('Before rebuilding this Word baseline, preserve team edits and comments. Regenerating the file replaces its contents.', small=True)

doc.save(OUT)
print(OUT.name)
print(json.dumps({'tasks': len(DATA['tasks']), 'roles': len(DATA['roles']), 'month_views': 7, 'editable_tables': len(doc.tables), 'file_bytes': OUT.stat().st_size}))
